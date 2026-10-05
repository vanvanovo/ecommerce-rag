# -*- coding: utf-8 -*-
"""
文档加载：把 PDF / Word / PPT / 图片 / Markdown / TXT 统一转换为 LangChain Document。

处理优先级（逐级降级，缺依赖不报错）：
  PDF   : MinerU(可选,保结构) → PyMuPDF 文本 + 内置表格识别 → 扫描页 OCR(RapidOCR)
  DOCX  : python-docx，按文档流顺序输出 段落 + 表格(Markdown)
  PPTX  : python-pptx，按位置排序输出文本框 + 表格
  图片   : RapidOCR 取文字；文字太少视为示意图 → BLIP 生成描述(可选)
  MD/TXT: 直接读取

元数据：source(父目录名) / sm_ty(产品类型) / tag(文件名关键词) / page_no / timestamp
"""
import os
from datetime import datetime
from typing import List

from langchain_core.documents import Document

from base.config import config
from base.logger import logger

# 目录名 → 产品类型
SM_TY_MAP = {"手机": "phone", "电脑": "laptop", "平板": "tablet"}
# 文件名关键词 → 标签
TAG_KW = [("价格", "价格"), ("参数", "参数"), ("保修", "保修"), ("退换", "退换货"),
          ("售后", "售后"), ("操作", "操作"), ("指南", "操作"), ("发票", "发票"),
          ("常见问题", "FAQ"), ("FAQ", "FAQ")]

SUPPORTED = {".pdf", ".md", ".txt", ".docx", ".pptx", ".png", ".jpg", ".jpeg", ".bmp"}


# ============================ OCR ============================
_OCR = None


def _get_ocr():
    """RapidOCR 引擎（单例）。"""
    global _OCR
    if _OCR is None:
        from rapidocr_onnxruntime import RapidOCR
        _OCR = RapidOCR()
    return _OCR


def _ocr(img):
    """img 支持 路径 或 numpy 数组；失败返回空串。"""
    try:
        result, _ = _get_ocr()(img)
        return "\n".join(line[1] for line in result) if result else ""
    except Exception as e:
        logger.warning(f"OCR 失败: {e}")
        return ""


def _ocr_pdf_page(page):
    """整页渲染后 OCR（给扫描页用）。"""
    import numpy as np
    import fitz
    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
    arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
    if pix.n == 4:
        arr = arr[:, :, :3]
    return _ocr(arr)


# ============================ MinerU（可选） ============================
def _mineru_markdown(pdf_path):
    """装了 mineru 就优先用；未装/失败返回 None，调用方自动降级。"""
    try:
        from mineru import MinerU  # noqa
    except Exception:
        return None
    try:
        out = MinerU()(pdf_path)
        if isinstance(out, dict):
            return out.get("md_content") or out.get("markdown")
        return str(out) if out else None
    except Exception as e:
        logger.warning(f"MinerU 解析失败，回退 PyMuPDF: {e}")
        return None


# ============================ PDF ============================
def _pdf_extract(pdf_path):
    md = _mineru_markdown(pdf_path)
    if md and md.strip():
        logger.info(f"MinerU 解析成功: {pdf_path}")
        return md

    import fitz
    doc = fitz.open(pdf_path)
    parts = []
    for page in doc:
        text = page.get_text().strip()
        # 内置表格识别（PyMuPDF >= 1.23 实验特性，失败忽略）
        try:
            for t in page.find_tables().tables:
                rows = [r for r in t.extract() if r]
                if rows:
                    parts.append("\n".join(
                        "| " + " | ".join((c or "").replace("\n", " ").strip() for c in row) + " |"
                        for row in rows))
        except Exception:
            pass
        if len(text) >= 20:
            parts.append(text)
        else:
            ocr_text = _ocr_pdf_page(page)     # 扫描页/图片页
            if ocr_text:
                parts.append(ocr_text)
    doc.close()
    return "\n\n".join(parts)


# ============================ DOCX ============================
def _docx_extract(path):
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    doc = DocxDocument(path)
    parts = []
    for child in doc.element.body.iterchildren():   # 按文档流顺序
        if child.tag == qn('w:p'):
            t = Paragraph(child, doc).text.strip()
            if t:
                parts.append(t)
        elif child.tag == qn('w:tbl'):
            rows = []
            for row in Table(child, doc).rows:
                rows.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
            if rows:
                parts.append("\n".join(rows))
    return "\n\n".join(parts)


# ============================ PPTX ============================
def _pptx_extract(path):
    from pptx import Presentation

    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"## 第 {i} 页")
        for sh in sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0)):   # 按位置排序
            if sh.has_text_frame and sh.text_frame.text.strip():
                parts.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    parts.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
    return "\n".join(parts)


# ============================ 图片 ============================
def _blip_describe(img_path):
    """示意图描述（可选）：仅当本地有 BLIP 模型时启用，避免联网卡住。"""
    model_dir = os.path.join(config.MODELS_DIR, "blip-image-captioning-base")
    if not os.path.isdir(model_dir):
        return ""
    try:
        from transformers import BlipProcessor, BlipForConditionalGeneration
        from PIL import Image
        proc = BlipProcessor.from_pretrained(model_dir)
        model = BlipForConditionalGeneration.from_pretrained(model_dir)
        img = Image.open(img_path).convert("RGB")
        out = model.generate(**proc(img, return_tensors="pt"), max_new_tokens=60)
        return proc.decode(out[0], skip_special_tokens=True)
    except Exception as e:
        logger.warning(f"BLIP 描述失败: {e}")
        return ""


def _image_extract(path):
    text = _ocr(path)
    if len(text.strip()) < 5:          # 没识别出文字 → 当示意图
        text = _blip_describe(path)
    return text


# ============================ 统一入口 ============================
def _extract_text(path, ext):
    if ext == ".pdf":
        return _pdf_extract(path)
    if ext == ".docx":
        return _docx_extract(path)
    if ext == ".pptx":
        return _pptx_extract(path)
    if ext in (".md", ".txt"):
        return open(path, encoding="utf-8", errors="ignore").read()
    if ext in (".png", ".jpg", ".jpeg", ".bmp"):
        return _image_extract(path)
    return ""


def _infer_tag(name):
    for kw, tag in TAG_KW:
        if kw in name:
            return tag
    return ""


def load_document_to_text(directory: str) -> List[Document]:
    """遍历目录，把支持的文件转成 Document（source=父目录名）。"""
    docs = []
    for root, _, files in os.walk(directory):
        for fn in files:
            ext = os.path.splitext(fn)[1].lower()
            if ext not in SUPPORTED:
                continue
            path = os.path.join(root, fn)
            source = os.path.basename(os.path.dirname(path))
            try:
                content = _extract_text(path, ext)
            except Exception as e:
                logger.error(f"加载失败 {path}: {e}")
                continue
            if content and content.strip():
                docs.append(Document(
                    page_content=content.strip(),
                    metadata={
                        "file_path": path,
                        "source": source,
                        "sm_ty": SM_TY_MAP.get(source, ""),
                        "tag": _infer_tag(fn),
                        "page_no": 0,
                        "timestamp": datetime.now().isoformat(),
                    },
                ))
                logger.info(f"加载成功: {path}（{len(content)} 字）")
            else:
                logger.warning(f"未提取到内容: {path}")
    return docs


if __name__ == "__main__":
    for d in load_document_to_text(config.DATA_DIR):
        print(d.metadata["source"], d.metadata["tag"], "|", d.page_content[:60].replace("\n", " "))
